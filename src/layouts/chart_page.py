from pptx.util import Inches, Pt

from src.components.text import add_title, add_bullets
from src.components.chart import (
    add_bar_chart,
    add_line_chart,
    add_pie_chart,
    add_waterfall,
    add_stacked_bar_chart,
    add_area_chart,
    add_scatter_chart,
    add_combo_chart,
)
from src.components._style import set_paragraph_text


CHART_FUNCTIONS = {
    "bar": add_bar_chart,
    "line": add_line_chart,
    "pie": add_pie_chart,
    "waterfall": add_waterfall,
    "stacked_bar": add_stacked_bar_chart,
    "area": add_area_chart,
    "scatter": add_scatter_chart,
    "combo": add_combo_chart,
}


class ChartPageLayout:
    def render(self, slide, theme, data):
        """チャート主体ページ: チャート + オプションでキーポイント + 出典"""
        title = data.get("title", "")
        subtitle = data.get("subtitle")
        chart_data = data.get("chart", {})
        key_points = data.get("key_points", [])
        source = data.get("source", "")

        if title:
            add_title(slide, theme, title, theme.margin_left, theme.margin_top, subtitle=subtitle)

        content_top = theme.content_area_top
        if subtitle:
            content_top += Inches(0.35)

        chart_type = chart_data.get("type", "bar")
        chart_func = CHART_FUNCTIONS.get(chart_type, add_bar_chart)
        unit = chart_data.get("unit", None)
        annotations = chart_data.get("annotations", [])

        extra_kwargs = {}
        if unit and chart_type in ("bar", "line", "stacked_bar", "area"):
            extra_kwargs["unit"] = unit
        if annotations and chart_type in ("bar", "line", "stacked_bar", "area", "combo"):
            extra_kwargs["annotations"] = annotations
        if chart_type == "stacked_bar" and chart_data.get("horizontal"):
            extra_kwargs["horizontal"] = True
        if chart_type == "area" and chart_data.get("stacked"):
            extra_kwargs["stacked"] = True
        if chart_type == "scatter":
            if chart_data.get("x_label"):
                extra_kwargs["x_label"] = chart_data["x_label"]
            if chart_data.get("y_label"):
                extra_kwargs["y_label"] = chart_data["y_label"]

        source_reserve = Inches(0.35) if source else 0
        available_height = theme.content_height - source_reserve
        if subtitle:
            available_height -= Inches(0.35)

        if key_points:
            chart_width = int(theme.content_width * 0.65)
            chart_func(
                slide, theme, chart_data["data"],
                theme.margin_left, content_top,
                width=chart_width, height=available_height,
                **extra_kwargs,
            )

            kp_left = theme.margin_left + chart_width + Inches(0.4)
            kp_width = int(theme.content_width * 0.30)
            add_bullets(
                slide, theme, key_points,
                kp_left, content_top + Inches(0.5),
                width=kp_width,
            )
        else:
            chart_func(
                slide, theme, chart_data["data"],
                theme.margin_left, content_top,
                width=theme.content_width, height=available_height,
                **extra_kwargs,
            )

        if source:
            source_top = content_top + available_height + Inches(0.05)
            src_box = slide.shapes.add_textbox(
                theme.margin_left, source_top,
                theme.content_width, Inches(0.3),
            )
            src_tf = src_box.text_frame
            src_tf.word_wrap = True
            set_paragraph_text(
                src_tf.paragraphs[0], f"出典: {source}",
                size=theme.font_size_footnote, color=theme.text_secondary,
                name=theme.font_body, italic=True,
            )
