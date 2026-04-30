"""出典 (source) 注記の共通描画ヘルパー。フッター上端の直上に小さく配置。

generate_pptx() がフッター付与前に各スライドの data.source / data.sources を
読み取って呼ぶ。chart_page など個別レイアウトは描画位置の予約だけ知っていればよい。
"""
from pptx.util import Inches

from src.components._style import set_paragraph_text


# フッターディバイダ直上に確保する注記領域の高さ (chart_page など領域計算で使う)
SOURCE_NOTE_HEIGHT = Inches(0.28)
SOURCE_NOTE_GAP = Inches(0.05)  # ディバイダとの隙間


def has_source(data: dict) -> bool:
    """source または sources が指定されているか。"""
    return bool(data.get("source")) or bool(data.get("sources"))


def render_source_note(slide, theme, data: dict) -> None:
    """data.source (str) または data.sources (list[str|{label,url}]) を1行に集約して描画。"""
    text = _format_source_text(data)
    if not text:
        return

    sw = theme.slide_width
    sh = theme.slide_height

    divider_top = sh - theme.margin_bottom + Inches(0.05)
    note_top = divider_top - SOURCE_NOTE_GAP - SOURCE_NOTE_HEIGHT

    box = slide.shapes.add_textbox(
        theme.margin_left,
        note_top,
        sw - theme.margin_left - theme.margin_right,
        SOURCE_NOTE_HEIGHT,
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_top = 0
    tf.margin_bottom = 0
    set_paragraph_text(
        tf.paragraphs[0],
        text,
        size=theme.font_size_footnote,
        color=theme.text_secondary,
        name=theme.font_body,
        italic=True,
    )


def _format_source_text(data: dict) -> str:
    if data.get("source"):
        return f"出典: {data['source']}"
    sources = data.get("sources")
    if not sources:
        return ""
    parts = []
    for s in sources:
        if isinstance(s, str):
            parts.append(s)
        elif isinstance(s, dict):
            label = s.get("label", "")
            url = s.get("url", "")
            if label and url:
                parts.append(f"{label} ({url})")
            elif label:
                parts.append(label)
            elif url:
                parts.append(url)
    if not parts:
        return ""
    return "出典: " + " / ".join(parts)
